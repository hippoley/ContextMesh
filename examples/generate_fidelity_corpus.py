from __future__ import annotations

import json
import shutil
import subprocess
import sys
import wave
from pathlib import Path


def run(cmd: list[str]) -> bool:
    try:
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return True
    except Exception:
        return False


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "./fidelity-corpus").resolve()
    out.mkdir(parents=True, exist_ok=True)

    from PIL import Image, ImageDraw
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from pptx import Presentation
    from pptx.util import Inches

    probes: dict[str, str] = {}

    p = out / "policy.txt"
    probes[p.name] = "NEEDLE-TXT-8742"
    p.write_text(f"The controlling marker is {probes[p.name]}.\n", encoding="utf-8")

    image = out / "visual.png"
    probes[image.name] = "NEEDLE-IMAGE-8742"
    im = Image.new("RGB", (1280, 720), "white")
    d = ImageDraw.Draw(im)
    d.text((120, 320), probes[image.name], fill="black")
    im.save(image)

    scan = out / "scan.pdf"
    probes[scan.name] = probes[image.name]
    im.save(scan, "PDF")

    ppt = out / "deck.pptx"
    probes[ppt.name] = "NEEDLE-PPT-8742"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = probes[ppt.name]
    slide.shapes.add_picture(str(image), Inches(1), Inches(2), width=Inches(6))
    prs.save(ppt)

    xlsx = out / "book.xlsx"
    probes[xlsx.name] = "NEEDLE-XLSX-8742"
    wb = Workbook()
    ws = wb.active
    ws.title = "Needles"
    ws.append(["marker", "a", "b", "formula"])
    ws.append([probes[xlsx.name], 1, 2, "=B2+C2"])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    ws.add_chart(chart, "F2")
    wb.save(xlsx)

    docx = out / "report.docx"
    probes[docx.name] = "NEEDLE-DOCX-8742"
    doc = Document()
    doc.add_heading("Fidelity", 1)
    doc.add_paragraph(probes[docx.name])
    doc.save(docx)

    audio = out / "speech.wav"
    probes[audio.name] = "needle audio eight seven four two"
    if shutil.which("espeak"):
        run(["espeak", "-w", str(audio), probes[audio.name]])
    else:
        with wave.open(str(audio), "wb") as w:
            w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000); w.writeframes(b"\x00\x00" * 16000)

    video = out / "clip.mp4"
    probes[video.name] = "NEEDLE-VIDEO-8742"
    video_frame = out / "video-frame.png"
    vim = Image.new("RGB", (1280, 720), "white")
    ImageDraw.Draw(vim).text((120, 320), probes[video.name], fill="black")
    vim.save(video_frame)
    if shutil.which("ffmpeg"):
        run(["ffmpeg", "-y", "-loop", "1", "-i", str(video_frame), "-i", str(audio), "-t", "2", "-shortest", "-pix_fmt", "yuv420p", str(video)])

    (out / "probes.json").write_text(json.dumps(probes, indent=2), encoding="utf-8")
    print(out)
    print(json.dumps(probes, indent=2))


if __name__ == "__main__":
    main()
