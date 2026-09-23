from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
import wave
from pathlib import Path

from contextmesh.ingest import ingest_paths
from contextmesh.store import FileContextStore


def build_assets(root: Path) -> list[Path]:
    from PIL import Image, ImageDraw
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document

    assets: list[Path] = []

    txt = root / "policy.txt"
    txt.write_text("NEEDLE-TXT-42: enterprise exception applies.", encoding="utf-8")
    assets.append(txt)

    image = root / "visual.png"
    im = Image.new("RGB", (800, 450), "white")
    ImageDraw.Draw(im).text((60, 180), "NEEDLE-IMAGE-42", fill="black")
    im.save(image)
    assets.append(image)

    scan = root / "scan.pdf"
    im.save(scan, "PDF")
    assets.append(scan)

    ppt = root / "deck.pptx"
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_picture(str(image), Inches(1), Inches(1), width=Inches(6))
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "NEEDLE-PPT-TEXT-42"
    prs.save(ppt)
    assets.append(ppt)

    xlsx = root / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Needles"
    ws.append(["key", "a", "b", "formula"])
    ws.append(["NEEDLE-XLSX-42", 1, 2, "=B2+C2"])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    ws.add_chart(chart, "F2")
    wb.save(xlsx)
    assets.append(xlsx)

    docx = root / "report.docx"
    doc = Document()
    doc.add_heading("Needle", 1)
    doc.add_paragraph("NEEDLE-DOCX-42 is authoritative.")
    doc.save(docx)
    assets.append(docx)

    wav = root / "silence.wav"
    with wave.open(str(wav), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(8000)
        w.writeframes(b"\x00\x00" * 8000)
    assets.append(wav)

    if shutil.which("ffmpeg"):
        video = root / "clip.mp4"
        subprocess.run([
            "ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=black:s=160x120:d=1",
            "-pix_fmt", "yuv420p", str(video),
        ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)
        assets.append(video)

    return assets


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="contextmesh-fidelity-") as td:
        root = Path(td)
        assets = build_assets(root)
        store = FileContextStore(root / "store")
        manifest = ingest_paths(assets, store, "fidelity-smoke")
        out = {
            "assets": len(manifest.assets),
            "required_blocks": manifest.required_blocks,
            "ingest_coverage": manifest.ingest_coverage,
            "semantic_coverage": manifest.semantic_coverage,
            "coverage_ready": manifest.coverage_ready,
            "unresolved_units": manifest.unresolved_units,
            "required_capabilities": manifest.required_capabilities,
            "modality_counts": manifest.modality_counts,
            "asset_reports": [x.model_dump() for x in manifest.asset_reports],
        }
        print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
