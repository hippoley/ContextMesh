from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path
from typing import Callable, Iterable

from .docling_adapter import ParsedDocument, ParsedItem, parse_with_docling
from .models import AssetCoverageReport, BlockKind, ContextBlock, CorpusManifest, Modality, SemanticStatus, SourceRef
from .store import FileContextStore
from .rich_ingest import parse_audio, parse_docx, parse_image, parse_pdf, parse_pptx, parse_video, parse_xlsx

TEXT_EXTS = {".txt", ".md", ".markdown", ".html", ".htm", ".xml", ".json", ".yaml", ".yml", ".py", ".js", ".ts", ".java", ".go", ".rs", ".sql"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".tiff", ".bmp"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".flac", ".aac", ".ogg"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".webm", ".avi"}
PDF_EXTS = {".pdf"}
PPT_EXTS = {".pptx"}
XLSX_EXTS = {".xlsx"}
DOCX_EXTS = {".docx"}


def _stable_id(*parts: str) -> str:
    return hashlib.sha256("\x1f".join(parts).encode()).hexdigest()[:20]


def _sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            chunk = f.read(chunk_size)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def _windows(text: str, size: int = 12000, overlap: int = 800) -> Iterable[tuple[int, int, str]]:
    if not text:
        return
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        yield start, end, text[start:end]
        if end == len(text):
            break
        start = max(start + 1, end - overlap)


def _modality_for_path(path: Path) -> Modality:
    ext = path.suffix.lower()
    if ext == ".csv" or ext in {".xls", ".xlsx", ".ods"}:
        return Modality.TABLE
    if ext in IMAGE_EXTS:
        return Modality.IMAGE
    if ext in AUDIO_EXTS:
        return Modality.AUDIO
    if ext in VIDEO_EXTS:
        return Modality.VIDEO
    return Modality.TEXT


def _make_root(corpus_id: str, path: Path, asset_id: str, modality: Modality) -> ContextBlock:
    return ContextBlock(
        id=_stable_id(asset_id, "root"),
        corpus_id=corpus_id,
        modality=modality,
        kind=BlockKind.ASSET,
        title=path.name,
        text="",
        source=SourceRef(asset_id=asset_id, path=str(path), locator={"asset": True}),
        depth=0,
        processable=False,
        metadata={"filename": path.name, "extension": path.suffix.lower()},
    )


def _put_sibling_chain(store: FileContextStore, blocks: list[ContextBlock]) -> None:
    for i, block in enumerate(blocks):
        block.prev_id = blocks[i - 1].id if i > 0 else None
        block.next_id = blocks[i + 1].id if i + 1 < len(blocks) else None
    store.put_blocks(blocks)


def _native_text_blocks(
    text: str,
    corpus_id: str,
    path: Path,
    asset_id: str,
    root: ContextBlock,
    window_chars: int,
    overlap_chars: int,
) -> tuple[list[ContextBlock], list[ContextBlock]]:
    """Create leaf blocks plus optional heading-based structural blocks.

    Markdown headings create addressable section parents. The original heading line
    stays in the leaf text, so structural indexing never replaces source content.
    """
    leafs: list[ContextBlock] = []
    structural: list[ContextBlock] = []
    is_markdown = path.suffix.lower() in {".md", ".markdown"}

    if not is_markdown:
        windows = list(_windows(text, window_chars, overlap_chars))
        for start, end, chunk in windows:
            leafs.append(
                ContextBlock(
                    id=_stable_id(asset_id, str(start), str(end)),
                    corpus_id=corpus_id,
                    modality=Modality.TEXT,
                    kind=BlockKind.CONTENT,
                    text=chunk,
                    source=SourceRef(asset_id=asset_id, path=str(path), locator={"char_start": start, "char_end": end}),
                    parent_id=root.id,
                    depth=1,
                    metadata={"parser": "native"},
                )
            )
        root.children_ids = [b.id for b in leafs]
        return leafs, structural

    heading_re = re.compile(r"^(#{1,6})\s+(.+?)\s*$", re.MULTILINE)
    matches = list(heading_re.finditer(text))
    if not matches:
        windows = list(_windows(text, window_chars, overlap_chars))
        for start, end, chunk in windows:
            leafs.append(
                ContextBlock(
                    id=_stable_id(asset_id, str(start), str(end)),
                    corpus_id=corpus_id,
                    modality=Modality.TEXT,
                    kind=BlockKind.CONTENT,
                    text=chunk,
                    source=SourceRef(asset_id=asset_id, path=str(path), locator={"char_start": start, "char_end": end}),
                    parent_id=root.id,
                    depth=1,
                    metadata={"parser": "native-markdown"},
                )
            )
        root.children_ids = [b.id for b in leafs]
        return leafs, structural

    # Preserve any preamble before the first heading as ordinary content.
    sections: list[tuple[int, int, int, str]] = []
    if matches[0].start() > 0:
        sections.append((0, matches[0].start(), 1, "Preamble"))
    for idx, match in enumerate(matches):
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        sections.append((match.start(), end, len(match.group(1)), match.group(2).strip()))

    stack: list[tuple[int, ContextBlock]] = []
    for sec_idx, (start, end, level, title) in enumerate(sections):
        while stack and stack[-1][0] >= level:
            stack.pop()
        parent = stack[-1][1] if stack else root
        section = ContextBlock(
            id=_stable_id(asset_id, "section", str(sec_idx), title),
            corpus_id=corpus_id,
            modality=Modality.TEXT,
            kind=BlockKind.SECTION,
            title=title,
            text="",
            source=SourceRef(asset_id=asset_id, path=str(path), locator={"char_start": start, "char_end": end, "heading_level": level}),
            parent_id=parent.id,
            depth=parent.depth + 1,
            processable=False,
            metadata={"parser": "native-markdown", "heading_level": level},
        )
        structural.append(section)
        parent.children_ids.append(section.id)
        stack.append((level, section))

        section_text = text[start:end]
        for local_start, local_end, chunk in _windows(section_text, window_chars, overlap_chars):
            leaf = ContextBlock(
                id=_stable_id(asset_id, "section", str(sec_idx), str(local_start), str(local_end)),
                corpus_id=corpus_id,
                modality=Modality.TEXT,
                kind=BlockKind.CONTENT,
                text=chunk,
                title=title,
                source=SourceRef(
                    asset_id=asset_id,
                    path=str(path),
                    locator={"char_start": start + local_start, "char_end": start + local_end, "section": title},
                ),
                parent_id=section.id,
                depth=section.depth + 1,
                metadata={"parser": "native-markdown", "heading_level": level},
            )
            leafs.append(leaf)
            section.children_ids.append(leaf.id)

    return leafs, structural


def _csv_blocks(
    path: Path,
    corpus_id: str,
    asset_id: str,
    root: ContextBlock,
    rows_per_block: int = 500,
) -> list[ContextBlock]:
    blocks: list[ContextBlock] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        try:
            header = next(reader)
        except StopIteration:
            return blocks
        batch: list[list[str]] = []
        row_start = 2
        for row_no, row in enumerate(reader, start=2):
            batch.append(row)
            if len(batch) >= rows_per_block:
                text = "\t".join(header) + "\n" + "\n".join("\t".join(r) for r in batch)
                blocks.append(
                    ContextBlock(
                        id=_stable_id(asset_id, "rows", str(row_start), str(row_no)),
                        corpus_id=corpus_id,
                        modality=Modality.TABLE,
                        kind=BlockKind.TABLE,
                        text=text,
                        source=SourceRef(asset_id=asset_id, path=str(path), locator={"row_start": row_start, "row_end": row_no}),
                        parent_id=root.id,
                        depth=1,
                        metadata={"parser": "native-csv", "columns": header},
                    )
                )
                batch = []
                row_start = row_no + 1
        if batch:
            row_end = row_start + len(batch) - 1
            text = "\t".join(header) + "\n" + "\n".join("\t".join(r) for r in batch)
            blocks.append(
                ContextBlock(
                    id=_stable_id(asset_id, "rows", str(row_start), str(row_end)),
                    corpus_id=corpus_id,
                    modality=Modality.TABLE,
                    kind=BlockKind.TABLE,
                    text=text,
                    source=SourceRef(asset_id=asset_id, path=str(path), locator={"row_start": row_start, "row_end": row_end}),
                    parent_id=root.id,
                    depth=1,
                    metadata={"parser": "native-csv", "columns": header},
                )
            )
    root.children_ids = [b.id for b in blocks]
    return blocks



def _native_pdf_blocks(path: Path, corpus_id: str, asset_id: str, root: ContextBlock, window_chars: int, overlap_chars: int) -> list[ContextBlock] | None:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return None
    blocks: list[ContextBlock] = []
    reader = PdfReader(str(path))
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            continue
        for start, end, chunk in _windows(text, window_chars, overlap_chars):
            blocks.append(ContextBlock(
                id=_stable_id(asset_id, "page", str(page_no), str(start), str(end)),
                corpus_id=corpus_id, modality=Modality.TEXT, kind=BlockKind.CONTENT, text=chunk,
                title=f"Page {page_no}", source=SourceRef(asset_id=asset_id, path=str(path), locator={"page": page_no, "page_char_start": start, "page_char_end": end}),
                parent_id=root.id, depth=1, metadata={"parser": "pypdf"},
            ))
    root.children_ids = [b.id for b in blocks]
    return blocks


def _native_pptx_blocks(path: Path, corpus_id: str, asset_id: str, root: ContextBlock, window_chars: int, overlap_chars: int) -> list[ContextBlock] | None:
    try:
        from pptx import Presentation  # type: ignore
    except Exception:
        return None
    blocks: list[ContextBlock] = []
    prs = Presentation(str(path))
    for slide_no, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                parts.append(str(text))
            if getattr(shape, "has_table", False):
                try:
                    rows = []
                    for row in shape.table.rows:
                        rows.append("\t".join(cell.text for cell in row.cells))
                    parts.append("\n".join(rows))
                except Exception:
                    pass
        text = "\n".join(x for x in parts if x).strip()
        if not text:
            text = "[Slide contains non-text visual content; install Docling for richer visual extraction.]"
        for start, end, chunk in _windows(text, window_chars, overlap_chars):
            blocks.append(ContextBlock(
                id=_stable_id(asset_id, "slide", str(slide_no), str(start), str(end)),
                corpus_id=corpus_id, modality=Modality.TEXT, kind=BlockKind.CONTENT, text=chunk, title=f"Slide {slide_no}",
                source=SourceRef(asset_id=asset_id, path=str(path), locator={"slide": slide_no, "slide_char_start": start, "slide_char_end": end}),
                parent_id=root.id, depth=1, metadata={"parser": "python-pptx"},
            ))
    root.children_ids = [b.id for b in blocks]
    return blocks


def _native_xlsx_blocks(path: Path, corpus_id: str, asset_id: str, root: ContextBlock, rows_per_block: int = 300) -> list[ContextBlock] | None:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return None
    blocks: list[ContextBlock] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            batch: list[list[str]] = []
            row_start = 1
            row_no = 0
            for row_no, values in enumerate(ws.iter_rows(values_only=True), start=1):
                batch.append(["" if v is None else str(v) for v in values])
                if len(batch) >= rows_per_block:
                    text = "\n".join("\t".join(r) for r in batch)
                    blocks.append(ContextBlock(
                        id=_stable_id(asset_id, "sheet", ws.title, str(row_start), str(row_no)), corpus_id=corpus_id,
                        modality=Modality.TABLE, kind=BlockKind.TABLE, text=text, title=ws.title,
                        source=SourceRef(asset_id=asset_id, path=str(path), locator={"sheet_name": ws.title, "row_start": row_start, "row_end": row_no}),
                        parent_id=root.id, depth=1, metadata={"parser": "openpyxl"},
                    ))
                    batch=[]; row_start=row_no+1
            if batch:
                row_end=row_start+len(batch)-1
                text="\n".join("\t".join(r) for r in batch)
                blocks.append(ContextBlock(
                    id=_stable_id(asset_id, "sheet", ws.title, str(row_start), str(row_end)), corpus_id=corpus_id,
                    modality=Modality.TABLE, kind=BlockKind.TABLE, text=text, title=ws.title,
                    source=SourceRef(asset_id=asset_id, path=str(path), locator={"sheet_name": ws.title, "row_start": row_start, "row_end": row_end}),
                    parent_id=root.id, depth=1, metadata={"parser": "openpyxl"},
                ))
    finally:
        wb.close()
    root.children_ids=[b.id for b in blocks]
    return blocks

def _docling_blocks(
    parsed: ParsedDocument,
    corpus_id: str,
    path: Path,
    asset_id: str,
    root: ContextBlock,
    window_chars: int,
    overlap_chars: int,
) -> tuple[list[ContextBlock], list[ContextBlock], list[str]]:
    leafs: list[ContextBlock] = []
    structural: list[ContextBlock] = []
    warnings: list[str] = []
    heading_stack: list[tuple[int, ContextBlock]] = []

    for idx, item in enumerate(parsed.items):
        if item.kind == BlockKind.SECTION:
            level = max(1, item.level)
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()
            parent = heading_stack[-1][1] if heading_stack else root
            section = ContextBlock(
                id=_stable_id(asset_id, "docling-section", str(idx), item.text),
                corpus_id=corpus_id,
                modality=Modality.TEXT,
                kind=BlockKind.SECTION,
                title=item.title or item.text,
                text="",
                source=SourceRef(asset_id=asset_id, path=str(path), locator=item.locator),
                parent_id=parent.id,
                depth=parent.depth + 1,
                processable=False,
                metadata={**parsed.metadata, **item.metadata},
            )
            parent.children_ids.append(section.id)
            structural.append(section)
            heading_stack.append((level, section))
            continue

        parent = heading_stack[-1][1] if heading_stack else root
        text = item.text or ""
        windows = list(_windows(text, window_chars, overlap_chars)) if text else [(0, 0, "")]
        for w_idx, (start, end, chunk) in enumerate(windows):
            unresolved_image = item.modality == Modality.IMAGE and not chunk and not item.metadata.get("media_path")
            block = ContextBlock(
                id=_stable_id(asset_id, "docling-item", str(idx), str(w_idx)),
                corpus_id=corpus_id,
                modality=item.modality,
                kind=item.kind,
                text=chunk,
                title=item.title,
                source=SourceRef(asset_id=asset_id, path=str(path), locator={**item.locator, "item_char_start": start, "item_char_end": end}),
                parent_id=parent.id,
                depth=parent.depth + 1,
                processable=not unresolved_image,
                semantic_status=SemanticStatus.UNRESOLVED if unresolved_image else SemanticStatus.READY,
                required_capabilities=["vision"] if item.modality == Modality.IMAGE else [],
                metadata={**parsed.metadata, **item.metadata},
            )
            leafs.append(block)
            parent.children_ids.append(block.id)
            if item.modality == Modality.IMAGE and not chunk:
                warnings.append(f"{path.name}: image block {block.id} has no generated caption/description; raw source is preserved")

    return leafs, structural, warnings


def ingest_paths(
    paths: list[Path],
    store: FileContextStore,
    corpus_id: str,
    window_chars: int = 12000,
    overlap_chars: int = 800,
    progress_callback: Callable[[dict], None] | None = None,
) -> CorpusManifest:
    """Ingest files into a model-agnostic, addressable corpus.

    v0.8 separates three contracts:
      * ingest coverage: did every submitted file produce an ingest result?
      * semantic coverage: are all known semantic channels addressable?
      * execution coverage: did the evaluator actually visit every required block?

    An unresolved visual/media channel never becomes a fake processable placeholder.
    """
    required_ids: list[str] = []
    structural_ids: list[str] = []
    root_ids: list[str] = []
    assets: list[str] = []
    total_chars = 0
    total_bytes = 0
    modality_counts: dict[str, int] = {}
    required_capabilities: set[str] = set()
    warnings: list[str] = []
    asset_reports: list[AssetCoverageReport] = []
    ingested_files = 0

    for file_index, path in enumerate(paths):
        if progress_callback:
            progress_callback({"phase": "file_started", "file_index": file_index, "total_files": len(paths), "path": str(path)})
        if not path.is_file():
            msg = f"skipped non-file input: {path}"
            warnings.append(msg)
            asset_reports.append(AssetCoverageReport(path=str(path), parser="none", status="failed", unresolved_units=1, warnings=[msg]))
            continue

        assets.append(str(path))
        file_bytes = path.stat().st_size
        file_sha256 = _sha256_file(path)
        total_bytes += file_bytes
        asset_id = _stable_id(corpus_id, str(path.resolve()))
        modality = _modality_for_path(path)
        root = _make_root(corpus_id, path, asset_id, modality)
        root_ids.append(root.id)
        ext = path.suffix.lower()
        structural: list[ContextBlock] = []
        unresolved_units = 0
        parser_name = "unknown"
        file_warnings: list[str] = []

        try:
            if ext == ".csv":
                leafs = _csv_blocks(path, corpus_id, asset_id, root)
                parser_name = "native-csv"
            elif ext in TEXT_EXTS:
                text = path.read_text(encoding="utf-8", errors="replace")
                leafs, structural = _native_text_blocks(text, corpus_id, path, asset_id, root, window_chars, overlap_chars)
                parser_name = "native-text"
            elif ext in IMAGE_EXTS:
                rich = parse_image(path, corpus_id, asset_id, root.id)
                leafs, parser_name, unresolved_units, file_warnings = rich.leafs, rich.parser, rich.unresolved_units, rich.warnings
            elif ext in AUDIO_EXTS:
                rich = parse_audio(path, corpus_id, asset_id, root.id)
                leafs, parser_name, unresolved_units, file_warnings = rich.leafs, rich.parser, rich.unresolved_units, rich.warnings
            elif ext in VIDEO_EXTS:
                rich = parse_video(path, corpus_id, asset_id, root.id)
                leafs, parser_name, unresolved_units, file_warnings = rich.leafs, rich.parser, rich.unresolved_units, rich.warnings
            else:
                # Prefer Docling when installed. Native fallbacks preserve a visual/raw
                # channel whenever possible so empty text extraction cannot masquerade
                # as semantic completeness.
                parsed = parse_with_docling(path)
                if parsed is not None:
                    leafs, structural, doc_warnings = _docling_blocks(parsed, corpus_id, path, asset_id, root, window_chars, overlap_chars)
                    file_warnings.extend(doc_warnings)
                    parser_name = "docling"
                    unresolved_units += sum(1 for b in leafs if b.semantic_status == SemanticStatus.UNRESOLVED)
                elif ext in PDF_EXTS:
                    rich = parse_pdf(path, corpus_id, asset_id, root.id, window_chars, overlap_chars)
                    if rich is None:
                        raise ValueError("PDF support requires Docling or pypdf")
                    leafs, parser_name, unresolved_units, file_warnings = rich.leafs, rich.parser, rich.unresolved_units, rich.warnings
                elif ext in PPT_EXTS:
                    rich = parse_pptx(path, corpus_id, asset_id, root.id, window_chars, overlap_chars)
                    if rich is None:
                        raise ValueError("PPTX support requires Docling or python-pptx")
                    leafs, parser_name, unresolved_units, file_warnings = rich.leafs, rich.parser, rich.unresolved_units, rich.warnings
                elif ext in XLSX_EXTS:
                    rich = parse_xlsx(path, corpus_id, asset_id, root.id)
                    if rich is None:
                        raise ValueError("XLSX support requires Docling or openpyxl")
                    leafs, parser_name, unresolved_units, file_warnings = rich.leafs, rich.parser, rich.unresolved_units, rich.warnings
                elif ext in DOCX_EXTS:
                    rich = parse_docx(path, corpus_id, asset_id, root.id, window_chars, overlap_chars)
                    if rich is None:
                        raise ValueError("DOCX support requires Docling or python-docx")
                    leafs, parser_name, unresolved_units, file_warnings = rich.leafs, rich.parser, rich.unresolved_units, rich.warnings
                else:
                    raise ValueError(
                        f"Unsupported {ext or 'file'} without Docling installed: {path.name}. "
                        "Install contextmesh[docling] for rich documents/media."
                    )

            # Never count an unresolved placeholder as a coverage unit. PARTIAL blocks
            # may be processed for the information they do contain, but the manifest
            # remains semantic-incomplete until their unresolved channel is fixed.
            processable_leafs = [b for b in leafs if b.processable and b.semantic_status != SemanticStatus.UNRESOLVED]
            unresolved_units += sum(1 for b in leafs if b.semantic_status == SemanticStatus.UNRESOLVED)
            partial_count = sum(1 for b in processable_leafs if b.semantic_status == SemanticStatus.PARTIAL)
            unresolved_units += partial_count

            root.children_ids = list(dict.fromkeys([*root.children_ids, *(b.id for b in leafs)]))
            _put_sibling_chain(store, leafs)
            store.put_blocks(structural)
            store.put_block(root)

            required_ids.extend(b.id for b in processable_leafs)
            modalities: dict[str, int] = {}
            for b in processable_leafs:
                key = b.modality.value
                modality_counts[key] = modality_counts.get(key, 0) + 1
                modalities[key] = modalities.get(key, 0) + 1
                required_capabilities.update(b.required_capabilities)
            structural_ids.extend([root.id] + [b.id for b in structural])
            total_chars += sum(len(b.text) for b in leafs)
            warnings.extend(file_warnings)
            status = "complete" if unresolved_units == 0 else "partial"
            report = AssetCoverageReport(
                path=str(path), parser=parser_name, bytes=file_bytes, source_sha256=file_sha256, status=status,
                resolved_units=len(processable_leafs), unresolved_units=unresolved_units,
                modalities=modalities, warnings=file_warnings,
            )
            asset_reports.append(report)
            ingested_files += 1

            if progress_callback:
                progress_callback({
                    "phase": "file_complete", "file_index": file_index, "total_files": len(paths), "path": str(path),
                    "required_blocks": len(processable_leafs), "modality": modality.value,
                    "semantic_coverage": report.semantic_coverage, "unresolved_units": unresolved_units,
                    "ingest_status": status,
                })
        except Exception as exc:
            msg = f"{path.name}: {type(exc).__name__}: {exc}"
            warnings.append(msg)
            asset_reports.append(AssetCoverageReport(
                path=str(path), parser=parser_name, bytes=file_bytes, source_sha256=file_sha256, status="failed", unresolved_units=1, warnings=[msg]
            ))
            store.put_block(root)
            if progress_callback:
                progress_callback({
                    "phase": "file_failed", "file_index": file_index, "total_files": len(paths), "path": str(path),
                    "error": msg, "semantic_coverage": 0.0, "unresolved_units": 1,
                })

    total_known_units = sum(r.resolved_units + r.unresolved_units for r in asset_reports)
    resolved_units = sum(r.resolved_units for r in asset_reports)
    unresolved_total = sum(r.unresolved_units for r in asset_reports)
    ingest_coverage = ingested_files / len(paths) if paths else 1.0
    semantic_coverage = resolved_units / total_known_units if total_known_units else (1.0 if ingest_coverage == 1.0 else 0.0)
    coverage_ready = ingest_coverage == 1.0 and unresolved_total == 0 and all(r.status == "complete" for r in asset_reports)

    manifest = CorpusManifest(
        corpus_id=corpus_id,
        assets=assets,
        root_block_ids=root_ids,
        structural_block_ids=structural_ids,
        block_ids=required_ids,
        required_block_ids=required_ids,
        total_blocks=len(required_ids),
        required_blocks=len(required_ids),
        total_chars=total_chars,
        total_bytes=total_bytes,
        modality_counts=modality_counts,
        required_capabilities=sorted(required_capabilities),
        ingest_warnings=warnings,
        asset_reports=asset_reports,
        ingest_coverage=ingest_coverage,
        semantic_coverage=semantic_coverage,
        unresolved_units=unresolved_total,
        coverage_ready=coverage_ready,
    )
    store.put_manifest(manifest)
    if progress_callback:
        progress_callback({
            "phase": "complete", "total_files": len(paths), "required_blocks": manifest.required_blocks,
            "corpus_id": corpus_id, "ingest_coverage": ingest_coverage, "semantic_coverage": semantic_coverage,
            "unresolved_units": unresolved_total, "coverage_ready": coverage_ready,
        })
    return manifest
