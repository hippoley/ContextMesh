from __future__ import annotations

import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

from .models import BlockKind, ContextBlock, Modality, SemanticStatus, SourceRef


@dataclass
class RichParseResult:
    leafs: list[ContextBlock] = field(default_factory=list)
    parser: str = "native"
    unresolved_units: int = 0
    warnings: list[str] = field(default_factory=list)


def _stable_id_fn(*parts: str) -> str:
    import hashlib
    return hashlib.sha256("\x1f".join(parts).encode()).hexdigest()[:20]


def _derived_dir(path: Path, asset_id: str) -> Path:
    d = path.parent / ".contextmesh_derived" / asset_id
    d.mkdir(parents=True, exist_ok=True)
    return d


def _text_windows(text: str, size: int, overlap: int) -> Iterable[tuple[int, int, str]]:
    if not text:
        return
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        yield start, end, text[start:end]
        if end == len(text):
            break
        start = max(start + 1, end - overlap)


def _image_block(
    *, corpus_id: str, asset_id: str, source_path: Path, media_path: Path, parent_id: str,
    title: str, locator: dict, block_id_parts: tuple[str, ...], parser: str, depth: int = 1,
) -> ContextBlock:
    return ContextBlock(
        id=_stable_id_fn(asset_id, *block_id_parts), corpus_id=corpus_id,
        modality=Modality.IMAGE, kind=BlockKind.IMAGE, text="", title=title,
        source=SourceRef(asset_id=asset_id, path=str(source_path), locator=locator),
        parent_id=parent_id, depth=depth, semantic_status=SemanticStatus.READY,
        required_capabilities=["vision"],
        metadata={"parser": parser, "media_path": str(media_path), "media_mime": _guess_image_mime(media_path)},
    )


def _guess_image_mime(path: Path) -> str:
    ext = path.suffix.lower()
    return {
        ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
        ".bmp": "image/bmp", ".tif": "image/tiff", ".tiff": "image/tiff",
    }.get(ext, "image/png")


def parse_image(path: Path, corpus_id: str, asset_id: str, root_id: str) -> RichParseResult:
    media_path = path
    # Normalize very large / awkward raster formats into a bounded JPEG payload for
    # OpenAI-compatible vision endpoints while preserving the original source path.
    try:
        if path.stat().st_size > 4_000_000 or path.suffix.lower() in {".tif", ".tiff", ".bmp"}:
            from PIL import Image  # type: ignore
            derived = _derived_dir(path, asset_id)
            dest = derived / "normalized.jpg"
            with Image.open(path) as im:
                im = im.convert("RGB")
                im.thumbnail((2200, 2200))
                im.save(dest, "JPEG", quality=82, optimize=True)
            media_path = dest
    except Exception:
        media_path = path
    return RichParseResult(
        leafs=[_image_block(
            corpus_id=corpus_id, asset_id=asset_id, source_path=path, media_path=media_path,
            parent_id=root_id, title=path.name, locator={"image": True}, block_id_parts=("image", "0"), parser="native-image",
        )],
        parser="native-image",
    )


def _render_pdf_pages(pdf_path: Path, out_dir: Path, prefix: str = "page") -> list[Path]:
    try:
        import fitz  # type: ignore
    except Exception:
        return []
    images: list[Path] = []
    doc = fitz.open(str(pdf_path))
    try:
        for idx, page in enumerate(doc, start=1):
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            dest = out_dir / f"{prefix}_{idx:04d}.png"
            pix.save(str(dest))
            images.append(dest)
    finally:
        doc.close()
    return images


def parse_pdf(
    path: Path, corpus_id: str, asset_id: str, root_id: str, window_chars: int, overlap_chars: int,
) -> RichParseResult | None:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return None
    result = RichParseResult(parser="pypdf+page-render")
    reader = PdfReader(str(path))
    derived = _derived_dir(path, asset_id)
    rendered = _render_pdf_pages(path, derived, "pdf_page")
    if not rendered:
        result.unresolved_units += len(reader.pages)
        result.warnings.append(f"{path.name}: PDF visual pages could not be rendered; visual semantics remain unresolved")

    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for start, end, chunk in _text_windows(text, window_chars, overlap_chars):
            result.leafs.append(ContextBlock(
                id=_stable_id_fn(asset_id, "pdf-text", str(page_no), str(start), str(end)), corpus_id=corpus_id,
                modality=Modality.TEXT, kind=BlockKind.CONTENT, text=chunk, title=f"Page {page_no} text",
                source=SourceRef(asset_id=asset_id, path=str(path), locator={"page": page_no, "channel": "text", "page_char_start": start, "page_char_end": end}),
                parent_id=root_id, depth=1, metadata={"parser": "pypdf"},
            ))
        if page_no <= len(rendered):
            result.leafs.append(_image_block(
                corpus_id=corpus_id, asset_id=asset_id, source_path=path, media_path=rendered[page_no - 1],
                parent_id=root_id, title=f"Page {page_no} visual",
                locator={"page": page_no, "channel": "visual"}, block_id_parts=("pdf-visual", str(page_no)), parser="pymupdf-render",
            ))
    return result


def _convert_office_to_pdf(path: Path, out_dir: Path) -> Path | None:
    try:
        proc = subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=90,
        )
    except (FileNotFoundError, subprocess.SubprocessError):
        return None
    candidate = out_dir / f"{path.stem}.pdf"
    if proc.returncode != 0 or not candidate.exists():
        return None
    return candidate


def parse_pptx(
    path: Path, corpus_id: str, asset_id: str, root_id: str, window_chars: int, overlap_chars: int,
) -> RichParseResult | None:
    try:
        from pptx import Presentation  # type: ignore
    except Exception:
        return None
    result = RichParseResult(parser="python-pptx+slide-render")
    prs = Presentation(str(path))
    derived = _derived_dir(path, asset_id)
    rendered: list[Path] = []
    pdf = _convert_office_to_pdf(path, derived)
    if pdf:
        rendered = _render_pdf_pages(pdf, derived, "slide")
    if len(rendered) < len(prs.slides):
        result.unresolved_units += len(prs.slides) - len(rendered)
        result.warnings.append(f"{path.name}: {len(prs.slides) - len(rendered)} slide visual(s) could not be rendered")

    for slide_no, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", None)
            if text:
                parts.append(str(text))
            if getattr(shape, "has_table", False):
                try:
                    parts.append("\n".join("\t".join(cell.text for cell in row.cells) for row in shape.table.rows))
                except Exception:
                    pass
        text = "\n".join(x for x in parts if x).strip()
        for start, end, chunk in _text_windows(text, window_chars, overlap_chars):
            result.leafs.append(ContextBlock(
                id=_stable_id_fn(asset_id, "slide-text", str(slide_no), str(start), str(end)), corpus_id=corpus_id,
                modality=Modality.TEXT, kind=BlockKind.CONTENT, text=chunk, title=f"Slide {slide_no} text",
                source=SourceRef(asset_id=asset_id, path=str(path), locator={"slide": slide_no, "channel": "text", "slide_char_start": start, "slide_char_end": end}),
                parent_id=root_id, depth=1, metadata={"parser": "python-pptx"},
            ))
        if slide_no <= len(rendered):
            result.leafs.append(_image_block(
                corpus_id=corpus_id, asset_id=asset_id, source_path=path, media_path=rendered[slide_no - 1],
                parent_id=root_id, title=f"Slide {slide_no} visual", locator={"slide": slide_no, "channel": "visual"},
                block_id_parts=("slide-visual", str(slide_no)), parser="libreoffice+pymupdf",
            ))
    return result


def _cell_repr(cell, cached_cell) -> str:
    value = cell.value
    if value is None and cached_cell.value is None:
        return ""
    if isinstance(value, str) and value.startswith("="):
        cached = cached_cell.value
        return f"{cell.coordinate}=FORMULA({value});cached={cached if cached is not None else '<none>'}"
    return f"{cell.coordinate}={value if value is not None else cached_cell.value}"


def _chart_description(chart) -> str:
    parts = [f"chart_type={type(chart).__name__}"]
    title = getattr(chart, "title", None)
    if title is not None:
        parts.append("title_present=true")
    for idx, series in enumerate(getattr(chart, "ser", []) or [], start=1):
        refs: list[str] = []
        for attr in ("val", "cat", "xVal", "yVal"):
            obj = getattr(series, attr, None)
            if obj is None:
                continue
            for refname in ("numRef", "strRef"):
                ref = getattr(obj, refname, None)
                formula = getattr(ref, "f", None) if ref is not None else None
                if formula:
                    refs.append(f"{attr}={formula}")
        parts.append(f"series_{idx}:" + (",".join(refs) if refs else "embedded/unknown"))
    return "\n".join(parts)


def parse_xlsx(path: Path, corpus_id: str, asset_id: str, root_id: str, rows_per_block: int = 300) -> RichParseResult | None:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return None
    result = RichParseResult(parser="openpyxl+formula-preserving")
    wb_formula = load_workbook(path, read_only=False, data_only=False)
    wb_values = load_workbook(path, read_only=False, data_only=True)
    derived = _derived_dir(path, asset_id)
    chart_count = 0
    image_count = 0
    try:
        for ws in wb_formula.worksheets:
            ws_values = wb_values[ws.title]
            rows: list[str] = []
            row_start = 1
            for row_no in range(1, ws.max_row + 1):
                cells: list[str] = []
                for col_no in range(1, ws.max_column + 1):
                    rep = _cell_repr(ws.cell(row_no, col_no), ws_values.cell(row_no, col_no))
                    if rep:
                        cells.append(rep)
                rows.append("\t".join(cells))
                if len(rows) >= rows_per_block:
                    result.leafs.append(ContextBlock(
                        id=_stable_id_fn(asset_id, "sheet", ws.title, str(row_start), str(row_no)), corpus_id=corpus_id,
                        modality=Modality.TABLE, kind=BlockKind.TABLE, text="\n".join(rows), title=ws.title,
                        source=SourceRef(asset_id=asset_id, path=str(path), locator={"sheet_name": ws.title, "row_start": row_start, "row_end": row_no}),
                        parent_id=root_id, depth=1,
                        metadata={"parser": "openpyxl", "sheet_state": ws.sheet_state, "merged_ranges": [str(x) for x in ws.merged_cells.ranges]},
                    ))
                    rows = []
                    row_start = row_no + 1
            if rows:
                row_end = row_start + len(rows) - 1
                result.leafs.append(ContextBlock(
                    id=_stable_id_fn(asset_id, "sheet", ws.title, str(row_start), str(row_end)), corpus_id=corpus_id,
                    modality=Modality.TABLE, kind=BlockKind.TABLE, text="\n".join(rows), title=ws.title,
                    source=SourceRef(asset_id=asset_id, path=str(path), locator={"sheet_name": ws.title, "row_start": row_start, "row_end": row_end}),
                    parent_id=root_id, depth=1,
                    metadata={"parser": "openpyxl", "sheet_state": ws.sheet_state, "merged_ranges": [str(x) for x in ws.merged_cells.ranges]},
                ))

            charts = list(getattr(ws, "_charts", []) or [])
            chart_count += len(charts)
            for idx, chart in enumerate(charts, start=1):
                result.leafs.append(ContextBlock(
                    id=_stable_id_fn(asset_id, "chart", ws.title, str(idx)), corpus_id=corpus_id,
                    modality=Modality.TABLE, kind=BlockKind.CONTENT, text=_chart_description(chart), title=f"{ws.title} chart {idx}",
                    source=SourceRef(asset_id=asset_id, path=str(path), locator={"sheet_name": ws.title, "chart": idx}),
                    parent_id=root_id, depth=1, semantic_status=SemanticStatus.READY,
                    metadata={"parser": "openpyxl", "chart_type": type(chart).__name__},
                ))

            images = list(getattr(ws, "_images", []) or [])
            image_count += len(images)
            for idx, image in enumerate(images, start=1):
                try:
                    raw = image._data()
                    ext = (getattr(image, "format", None) or "png").lower()
                    dest = derived / f"sheet_{ws.title}_image_{idx}.{ext}"
                    dest.write_bytes(raw)
                    result.leafs.append(_image_block(
                        corpus_id=corpus_id, asset_id=asset_id, source_path=path, media_path=dest, parent_id=root_id,
                        title=f"{ws.title} image {idx}", locator={"sheet_name": ws.title, "image": idx},
                        block_id_parts=("sheet-image", ws.title, str(idx)), parser="openpyxl-image",
                    ))
                except Exception:
                    result.unresolved_units += 1
                    result.warnings.append(f"{path.name}: could not materialize image {idx} in sheet {ws.title}")
    finally:
        wb_formula.close(); wb_values.close()

    # When a workbook contains charts/images, capture the rendered workbook too.
    # This preserves visual emphasis/layout that formula refs alone cannot certify.
    if chart_count or image_count:
        pdf = _convert_office_to_pdf(path, derived)
        rendered = _render_pdf_pages(pdf, derived, "workbook_page") if pdf else []
        if rendered:
            for page_no, img in enumerate(rendered, start=1):
                result.leafs.append(_image_block(
                    corpus_id=corpus_id, asset_id=asset_id, source_path=path, media_path=img, parent_id=root_id,
                    title=f"Workbook visual page {page_no}", locator={"workbook_page": page_no, "channel": "visual"},
                    block_id_parts=("workbook-visual", str(page_no)), parser="libreoffice+pymupdf",
                ))
        else:
            result.unresolved_units += chart_count + image_count
            result.warnings.append(
                f"{path.name}: workbook contains {chart_count} chart(s)/{image_count} image(s) but visual rendering failed"
            )
    return result


def parse_docx(path: Path, corpus_id: str, asset_id: str, root_id: str, window_chars: int, overlap_chars: int) -> RichParseResult | None:
    try:
        from docx import Document  # type: ignore
    except Exception:
        return None
    result = RichParseResult(parser="python-docx+page-render")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text)
    for start, end, chunk in _text_windows(text, window_chars, overlap_chars):
        result.leafs.append(ContextBlock(
            id=_stable_id_fn(asset_id, "docx-text", str(start), str(end)), corpus_id=corpus_id,
            modality=Modality.TEXT, kind=BlockKind.CONTENT, text=chunk, title=path.name,
            source=SourceRef(asset_id=asset_id, path=str(path), locator={"channel": "text", "char_start": start, "char_end": end}),
            parent_id=root_id, depth=1, metadata={"parser": "python-docx"},
        ))
    for t_idx, table in enumerate(doc.tables, start=1):
        rows = ["\t".join(cell.text for cell in row.cells) for row in table.rows]
        result.leafs.append(ContextBlock(
            id=_stable_id_fn(asset_id, "docx-table", str(t_idx)), corpus_id=corpus_id,
            modality=Modality.TABLE, kind=BlockKind.TABLE, text="\n".join(rows), title=f"Table {t_idx}",
            source=SourceRef(asset_id=asset_id, path=str(path), locator={"table": t_idx}),
            parent_id=root_id, depth=1, metadata={"parser": "python-docx"},
        ))
    derived = _derived_dir(path, asset_id)
    pdf = _convert_office_to_pdf(path, derived)
    rendered = _render_pdf_pages(pdf, derived, "docx_page") if pdf else []
    if rendered:
        for page_no, img in enumerate(rendered, start=1):
            result.leafs.append(_image_block(
                corpus_id=corpus_id, asset_id=asset_id, source_path=path, media_path=img, parent_id=root_id,
                title=f"Page {page_no} visual", locator={"page": page_no, "channel": "visual"},
                block_id_parts=("docx-visual", str(page_no)), parser="libreoffice+pymupdf",
            ))
    else:
        # Text/tables are still usable, but headers, floating shapes, footers and
        # layout-only semantics cannot be certified by the lightweight fallback.
        result.unresolved_units += 1
        result.warnings.append(f"{path.name}: DOCX visual/layout channel could not be rendered")
    return result


def _ffprobe_duration(path: Path) -> float | None:
    try:
        proc = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1:nk=1", str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=30,
        )
        if proc.returncode == 0:
            return float(proc.stdout.strip())
    except Exception:
        pass
    return None


def _extract_audio_segment(path: Path, dest: Path, start: float, duration: float) -> bool:
    try:
        proc = subprocess.run([
            "ffmpeg", "-y", "-ss", f"{start:.3f}", "-t", f"{duration:.3f}", "-i", str(path),
            "-vn", "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le", str(dest),
        ], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=max(30, int(duration * 2)))
        return proc.returncode == 0 and dest.exists() and dest.stat().st_size > 0
    except Exception:
        return False


def parse_audio(path: Path, corpus_id: str, asset_id: str, root_id: str, segment_seconds: int = 60) -> RichParseResult:
    duration = _ffprobe_duration(path)
    result = RichParseResult(parser="ffmpeg-audio-segments" if duration else "native-audio")
    if not duration or duration <= segment_seconds:
        media = path
        # Normalize unsupported chat-audio formats to WAV when ffmpeg is available.
        if path.suffix.lower() not in {".wav", ".mp3"} and duration:
            derived = _derived_dir(path, asset_id)
            dest = derived / "audio_0000.wav"
            if _extract_audio_segment(path, dest, 0.0, duration):
                media = dest
        result.leafs.append(ContextBlock(
            id=_stable_id_fn(asset_id, "audio", "0"), corpus_id=corpus_id,
            modality=Modality.AUDIO, kind=BlockKind.AUDIO, text="", title=path.name,
            source=SourceRef(asset_id=asset_id, path=str(path), locator={"start_time": 0.0, "end_time": duration}),
            parent_id=root_id, depth=1, semantic_status=SemanticStatus.READY,
            required_capabilities=["audio"], metadata={"parser": result.parser, "media_path": str(media), "duration_seconds": duration},
        ))
        return result

    derived = _derived_dir(path, asset_id)
    start = 0.0
    idx = 0
    while start < duration:
        end = min(duration, start + segment_seconds)
        dest = derived / f"audio_{idx:04d}.wav"
        if _extract_audio_segment(path, dest, start, end - start):
            result.leafs.append(ContextBlock(
                id=_stable_id_fn(asset_id, "audio", str(idx)), corpus_id=corpus_id,
                modality=Modality.AUDIO, kind=BlockKind.AUDIO, text="", title=f"{path.name} {start:.1f}-{end:.1f}s",
                source=SourceRef(asset_id=asset_id, path=str(path), locator={"start_time": start, "end_time": end}),
                parent_id=root_id, depth=1, semantic_status=SemanticStatus.READY,
                required_capabilities=["audio"], metadata={"parser": "ffmpeg-audio-segments", "media_path": str(dest), "duration_seconds": duration},
            ))
        else:
            result.unresolved_units += 1
            result.warnings.append(f"{path.name}: audio segment {start:.1f}-{end:.1f}s could not be materialized")
        idx += 1
        start = end
    return result


def _extract_video_segment(path: Path, dest: Path, start: float, duration: float) -> bool:
    try:
        proc = subprocess.run([
            "ffmpeg", "-y", "-ss", f"{start:.3f}", "-t", f"{duration:.3f}", "-i", str(path),
            "-c", "copy", str(dest),
        ], stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=max(30, int(duration * 2)))
        return proc.returncode == 0 and dest.exists() and dest.stat().st_size > 0
    except Exception:
        return False


def parse_video(path: Path, corpus_id: str, asset_id: str, root_id: str, segment_seconds: int = 60) -> RichParseResult:
    duration = _ffprobe_duration(path)
    result = RichParseResult(parser="ffmpeg-video-segments")
    if not duration or duration <= 0:
        result.unresolved_units = 1
        result.warnings.append(f"{path.name}: video duration could not be determined")
        return result
    derived = _derived_dir(path, asset_id)
    start = 0.0
    idx = 0
    while start < duration:
        end = min(duration, start + segment_seconds)
        dest = derived / f"video_{idx:04d}.mp4"
        if _extract_video_segment(path, dest, start, end - start):
            result.leafs.append(ContextBlock(
                id=_stable_id_fn(asset_id, "video", str(idx)), corpus_id=corpus_id,
                modality=Modality.VIDEO, kind=BlockKind.VIDEO, text="", title=f"{path.name} {start:.1f}-{end:.1f}s",
                source=SourceRef(asset_id=asset_id, path=str(path), locator={"start_time": start, "end_time": end}),
                parent_id=root_id, depth=1, semantic_status=SemanticStatus.READY,
                required_capabilities=["video"], metadata={"parser": "ffmpeg-video-segments", "media_path": str(dest), "duration_seconds": duration},
            ))
        else:
            result.unresolved_units += 1
            result.warnings.append(f"{path.name}: video segment {start:.1f}-{end:.1f}s could not be materialized")
        idx += 1
        start = end
    return result

