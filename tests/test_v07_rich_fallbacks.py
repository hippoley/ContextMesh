from pathlib import Path

from contextmesh.explorer import build_explorer_groups
from contextmesh.ingest import ingest_paths
from contextmesh.store import FileContextStore


def test_xlsx_sheet_explorer(tmp_path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["quarter", "value"])
    ws.append(["Q1", 10])
    ws.append(["Q2", 20])
    path = tmp_path / "book.xlsx"
    wb.save(path)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([path], store, "corp_xlsx")
    assert m.required_blocks >= 1
    groups = build_explorer_groups(store, m.corpus_id)
    assert any(g.kind == "sheet" and "Revenue" in g.label for g in groups)


def test_pptx_slide_explorer(tmp_path: Path):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "ContextMesh"
    slide.placeholders[1].text = "Full coverage"
    path = tmp_path / "deck.pptx"
    prs.save(path)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([path], store, "corp_pptx")
    groups = build_explorer_groups(store, m.corpus_id)
    assert any(g.kind == "slide" and g.label == "Slide 1" for g in groups)


def test_pdf_page_explorer(tmp_path: Path):
    # Build a tiny PDF with PyMuPDF; pypdf is used by the ingest fallback.
    import fitz
    path = tmp_path / "paper.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Full coverage PDF page")
    doc.save(path)
    doc.close()

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([path], store, "corp_pdf")
    groups = build_explorer_groups(store, m.corpus_id)
    assert any(g.kind == "page" and g.label == "Page 1" for g in groups)
