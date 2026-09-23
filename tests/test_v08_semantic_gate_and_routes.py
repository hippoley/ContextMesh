from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from contextmesh.ingest import ingest_paths
from contextmesh.judges import HeuristicJudge, OpenAICompatibleJudge
from contextmesh.models import Modality, ModelRoute
from contextmesh.runtime import ProgressiveEvaluator
from contextmesh.store import FileContextStore


def test_semantic_gate_blocks_score_even_after_all_addressable_text_is_visited(tmp_path: Path, monkeypatch):
    from docx import Document
    import contextmesh.rich_ingest as rich

    p = tmp_path / "layout.docx"
    doc = Document()
    doc.add_paragraph("addressable text survives")
    doc.save(p)

    # Force the lightweight fallback to lose the visual/layout channel.
    monkeypatch.setattr(rich, "_convert_office_to_pdf", lambda *args, **kwargs: None)
    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([p], store, "corp_partial")
    assert m.required_blocks >= 1
    assert m.ingest_coverage == 1.0
    assert m.semantic_coverage < 1.0
    assert m.coverage_ready is False

    result = ProgressiveEvaluator(store, HeuristicJudge()).evaluate(
        m.corpus_id, "addressable", "candidate"
    )
    assert result.coverage == 1.0  # execution read every available block
    assert result.ingest_ready is False
    assert result.complete is False
    assert result.score is None
    assert "semantic coverage is incomplete" in result.rationale.lower()


def test_model_route_capabilities_round_trip(tmp_path: Path):
    store = FileContextStore(tmp_path / "store")
    route = ModelRoute(
        id="vision-local", label="Vision local", base_url="http://localhost:8000/v1", model="vision-model",
        capabilities=["text", "table", "vision"],
    )
    store.upsert_model_route(route)
    loaded = store.list_model_routes()[0]
    assert loaded.capabilities == ["text", "table", "vision"]


def test_openai_compatible_judge_respects_declared_vision_capability(tmp_path: Path):
    from PIL import Image
    from contextmesh.models import BlockKind, ContextBlock, SourceRef

    p = tmp_path / "x.png"
    Image.new("RGB", (8, 8), "white").save(p)
    block = ContextBlock(
        id="b", corpus_id="c", modality=Modality.IMAGE, kind=BlockKind.IMAGE,
        source=SourceRef(asset_id="a", path=str(p)), metadata={"media_path": str(p)}, required_capabilities=["vision"],
    )
    text_only = OpenAICompatibleJudge("m", "http://example.invalid/v1", capabilities=["text", "table"])
    vision = OpenAICompatibleJudge("m", "http://example.invalid/v1", capabilities=["text", "vision"])
    assert text_only.can_inspect(block) is False
    assert vision.can_inspect(block) is True


@pytest.mark.skipif(shutil.which("libreoffice") is None, reason="LibreOffice required for workbook visual fallback")
def test_xlsx_chart_gets_structured_and_visual_coverage(tmp_path: Path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    p = tmp_path / "chart.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Quarter", "Revenue"])
    ws.append(["Q1", 10])
    ws.append(["Q2", 20])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
    ws.add_chart(chart, "D2")
    wb.save(p)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([p], store, "corp_chart")
    blocks = [store.get_block(m.corpus_id, x) for x in m.required_block_ids]
    assert any("chart_type=" in b.text for b in blocks)
    assert any(b.modality == Modality.IMAGE and b.source.locator.get("workbook_page") for b in blocks)
    assert m.coverage_ready is True
