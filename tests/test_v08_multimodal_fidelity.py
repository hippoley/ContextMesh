from __future__ import annotations

import shutil
import subprocess
import wave
from pathlib import Path

import pytest

from contextmesh.ingest import ingest_paths
from contextmesh.judges import HeuristicJudge
from contextmesh.models import Modality
from contextmesh.reader import CorpusReader
from contextmesh.runtime import ProgressiveEvaluator
from contextmesh.store import FileContextStore


class VisionSmokeJudge(HeuristicJudge):
    def can_inspect(self, block):
        if block.modality == Modality.IMAGE:
            media = block.metadata.get("media_path")
            return bool(media and Path(media).is_file())
        return super().can_inspect(block)

    def inspect(self, question, answer, block, notes):
        if block.modality == Modality.IMAGE:
            return f"{block.id} [image]: visual payload inspected", True
        return super().inspect(question, answer, block, notes)


def test_scan_pdf_becomes_visual_coverage_not_zero_blocks(tmp_path: Path):
    import fitz
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (600, 200), "white")
    ImageDraw.Draw(img).text((20, 80), "needle-vision-42", fill="black")
    img_path = tmp_path / "scan.png"
    img.save(img_path)

    pdf_path = tmp_path / "scan.pdf"
    doc = fitz.open()
    page = doc.new_page(width=600, height=200)
    page.insert_image(page.rect, filename=str(img_path))
    doc.save(pdf_path)
    doc.close()

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([pdf_path], store, "corp_scan")
    blocks = [store.get_block(m.corpus_id, x) for x in m.required_block_ids]

    assert m.required_blocks >= 1
    assert any(b.modality == Modality.IMAGE for b in blocks)
    assert m.coverage_ready is True
    assert m.semantic_coverage == 1.0

    # A text-only route must not pretend it read the scanned page.
    blocked = ProgressiveEvaluator(store, HeuristicJudge()).evaluate("corp_scan", "needle", "answer")
    assert blocked.score is None
    assert blocked.complete is False
    assert blocked.failed_blocks >= 1

    # A vision-capable route can complete the same corpus.
    ok = ProgressiveEvaluator(store, VisionSmokeJudge()).evaluate("corp_scan", "needle", "answer")
    assert ok.complete is True
    assert ok.coverage == 1.0
    assert ok.score is not None


def test_direct_image_requires_vision_capability(tmp_path: Path):
    from PIL import Image

    image_path = tmp_path / "diagram.png"
    Image.new("RGB", (40, 40), "white").save(image_path)
    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([image_path], store, "corp_image")
    block = store.get_block(m.corpus_id, m.required_block_ids[0])
    assert block.modality == Modality.IMAGE
    assert block.metadata.get("media_path") == str(image_path)
    assert m.coverage_ready is True

    r = ProgressiveEvaluator(store, HeuristicJudge()).evaluate("corp_image", "diagram", "answer")
    assert r.complete is False
    assert r.score is None
    assert r.failed_blocks == 1


def test_xlsx_preserves_formula_text_even_without_cached_result(tmp_path: Path):
    from openpyxl import Workbook

    p = tmp_path / "formula.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Calc"
    ws["A2"] = 1
    ws["B2"] = 2
    ws["C2"] = "=A2+B2"
    wb.save(p)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([p], store, "corp_formula")
    text = "\n".join(store.get_block(m.corpus_id, x).text for x in m.required_block_ids)
    assert "C2=FORMULA(=A2+B2)" in text
    assert m.coverage_ready is True


@pytest.mark.skipif(shutil.which("libreoffice") is None, reason="LibreOffice required for whole-slide visual fallback")
def test_image_only_pptx_has_real_slide_visual_not_placeholder(tmp_path: Path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    img_path = tmp_path / "slide-image.png"
    Image.new("RGB", (320, 180), "white").save(img_path)
    ppt_path = tmp_path / "visual.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(img_path), Inches(1), Inches(1), width=Inches(4))
    prs.save(ppt_path)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([ppt_path], store, "corp_ppt_visual")
    blocks = [store.get_block(m.corpus_id, x) for x in m.required_block_ids]
    assert any(b.modality == Modality.IMAGE and b.source.locator.get("slide") == 1 for b in blocks)
    assert all("install Docling" not in (b.text or "") for b in blocks)
    assert m.coverage_ready is True


def test_docx_text_is_addressable_and_layout_channel_is_accounted(tmp_path: Path):
    from docx import Document

    p = tmp_path / "report.docx"
    doc = Document()
    doc.add_heading("Policy", 1)
    doc.add_paragraph("The critical exception is forty-two.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "key"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "limit"
    table.cell(1, 1).text = "42"
    doc.save(p)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([p], store, "corp_docx")
    blocks = [store.get_block(m.corpus_id, x) for x in m.required_block_ids]
    joined = "\n".join(b.text for b in blocks)
    assert "critical exception" in joined
    assert "limit" in joined
    assert m.asset_reports[0].parser.startswith("python-docx")
    # On systems with LibreOffice/PyMuPDF the visual/layout channel becomes complete;
    # otherwise the manifest must be partial instead of lying about full fidelity.
    if m.coverage_ready:
        assert any(b.modality == Modality.IMAGE for b in blocks)
    else:
        assert m.unresolved_units > 0
        assert m.semantic_coverage < 1.0


def test_audio_is_addressable_but_text_only_judge_is_blocked(tmp_path: Path):
    p = tmp_path / "tone.wav"
    with wave.open(str(p), "wb") as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(8000)
        w.writeframes(b"\x00\x00" * 800)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([p], store, "corp_audio")
    block = store.get_block(m.corpus_id, m.required_block_ids[0])
    assert block.modality == Modality.AUDIO
    assert "audio" in block.required_capabilities
    assert m.coverage_ready is True
    r = ProgressiveEvaluator(store, HeuristicJudge()).evaluate("corp_audio", "what was said", "answer")
    assert r.score is None
    assert r.complete is False


@pytest.mark.skipif(shutil.which("ffmpeg") is None or shutil.which("ffprobe") is None, reason="ffmpeg required")
def test_video_is_split_into_timeline_blocks_and_cannot_fake_completion(tmp_path: Path):
    p = tmp_path / "clip.mp4"
    subprocess.run([
        "ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=black:s=160x120:d=2",
        "-pix_fmt", "yuv420p", str(p)
    ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)

    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([p], store, "corp_video")
    blocks = [store.get_block(m.corpus_id, x) for x in m.required_block_ids]
    assert blocks and all(b.modality == Modality.VIDEO for b in blocks)
    assert all("start_time" in b.source.locator for b in blocks)
    r = ProgressiveEvaluator(store, HeuristicJudge()).evaluate("corp_video", "what happened", "answer")
    assert r.score is None
    assert r.complete is False

@pytest.mark.parametrize(
    "suffix,content",
    [
        (".txt", "needle plain text 42"),
        (".md", "# Heading\nneedle markdown 42"),
        (".html", "<html><body>needle html 42</body></html>"),
        (".json", '{"needle": 42}'),
        (".xml", "<root><needle>42</needle></root>"),
        (".yaml", "needle: 42\n"),
    ],
)
def test_text_family_formats_are_addressable(tmp_path: Path, suffix: str, content: str):
    p = tmp_path / f"sample{suffix}"
    p.write_text(content, encoding="utf-8")
    store = FileContextStore(tmp_path / "store")
    m = ingest_paths([p], store, f"corp_{suffix[1:]}")
    assert m.coverage_ready is True
    assert m.required_blocks >= 1
    text = "\n".join(store.get_block(m.corpus_id, x).text for x in m.required_block_ids)
    assert "needle" in text
